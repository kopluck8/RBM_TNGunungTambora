"""Ubah format_laporan_rbm_bulanan.docx (dari kantor) menjadi template berplaceholder.

Pakai: python build_template.py <format.docx> <keluaran.docx>
Semua placeholder ditulis dalam SATU run agar mudah diganti oleh Apps Script.
"""
import copy
import sys

import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

src, out = sys.argv[1], sys.argv[2]
d = docx.Document(src)
body = d.element.body
kids = list(body.iterchildren())
P = lambda i: Paragraph(kids[i], d)


def set_text(p, text, italic=None):
    """Isi paragraf dengan satu run (format run pertama dipertahankan)."""
    runs = p.runs
    if runs:
        r0 = runs[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        r0.text = text
    else:
        r0 = p.add_run(text)
        r0.font.name = 'Arial'
        r0.font.size = docx.shared.Pt(11)
    if italic is not None:
        r0.italic = italic
    return r0


def clone_after(p, text):
    new = copy.deepcopy(p._p)
    p._p.addnext(new)
    np_ = Paragraph(new, p._parent)
    set_text(np_, text)
    return np_


def plain_after(p, text):
    """Paragraf polos (tanpa tebal/garis bawah) setelah p."""
    np_ = clone_after(p, text)
    r = np_.runs[0]
    r.bold = False
    r.underline = False
    r.italic = False
    np_.alignment = 0
    return np_


# ---------- COVER ----------
p_resor = P(5)
set_text(p_resor, '{{RESOR_KAPITAL}}')
clone_after(p_resor, '{{SPTN_KAPITAL}}')
body.remove(kids[18])  # satu baris kosong dibuang agar cover tetap satu halaman
set_text(P(20), 'PERIODE {{PERIODE_KAPITAL}}')
set_text(P(27), '{{TEMPAT_SPTN}}, {{PERIODE}}')

# ---------- LEMBAR PENGESAHAN ----------
info, ttd = d.tables[2], d.tables[3]

def isi_sel(cell, baris):
    """Ganti seluruh isi sel: tabel bersarang dibuang, paragraf diisi sesuai daftar."""
    for t in cell._tc.findall(qn('w:tbl')):
        cell._tc.remove(t)
    paras = cell.paragraphs
    while len(paras) < len(baris):
        paras[-1]._p.addnext(copy.deepcopy(paras[-1]._p))
        paras = cell.paragraphs
    for p, teks in zip(paras, baris):
        if isinstance(teks, list):  # [(teks, italic), ...]
            set_text(p, teks[0][0])
            base = p.runs[0]
            for t, it in teks[1:]:
                nr = copy.deepcopy(base._r)
                base._r.getparent().append(nr)
                from docx.text.run import Run
                rr = Run(nr, p)
                rr.text = t
                rr.italic = it
        else:
            set_text(p, teks)
    for p in paras[len(baris):]:
        p._p.getparent().remove(p._p)

# info: Judul / Dasar / Waktu / Lokasi
isi_sel(info.cell(0, 2), [[('Laporan Kegiatan ', False), ('Resort Based Management (RBM)', True), (' {{RESOR}}', False)]])
isi_sel(info.cell(1, 2), ['{{DASAR}}'])
isi_sel(info.cell(2, 2), ['{{PERIODE}}'])
isi_sel(info.cell(3, 2), ['{{RESOR}}, {{SPTN}}'])

# tanda tangan
isi_sel(ttd.cell(0, 0), ['', 'Menilai,', '{{JABATAN_SPTN}}'])
isi_sel(ttd.cell(0, 2), ['', '{{TEMPAT_TTD}}, {{TANGGAL_TTD}}', '{{JABATAN_KR}}'])
isi_sel(ttd.cell(1, 0), ['', '', '', '', '{{NAMA_SPTN}}', 'NIP. {{NIP_SPTN}}'])
isi_sel(ttd.cell(1, 2), ['', '', '', '', '{{NAMA_KR}}', 'NIP. {{NIP_KR}}'])
isi_sel(ttd.cell(2, 0), ['Mengetahui,', '{{JABATAN_BALAI}}', '', '', '', '', '{{NAMA_BALAI}}', 'NIP. {{NIP_BALAI}}'])
for row in ttd.rows:
    for c in row.cells:
        for p in c.paragraphs:
            p.alignment = 1  # tengah

# tabel tanda tangan tanpa garis
tblPr = ttd._tbl.tblPr
st = tblPr.find(qn('w:tblStyle'))
if st is not None:
    tblPr.remove(st)
b = OxmlElement('w:tblBorders')
for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    e = OxmlElement('w:' + side)
    e.set(qn('w:val'), 'nil')
    b.append(e)
tblPr.append(b)

# ---------- ISI & LAMPIRAN ----------
kids = list(body.iterchildren())
p_isi = next(Paragraph(k, d) for k in kids if k.tag == qn('w:p') and Paragraph(k, d).text.startswith('Isi Laporan'))
p_lamp = next(Paragraph(k, d) for k in kids if k.tag == qn('w:p') and Paragraph(k, d).text.startswith('Lampiran'))
set_text(p_isi, 'LAPORAN HASIL KEGIATAN RBM BULAN {{PERIODE_KAPITAL}}')
p_isi.alignment = 1
plain_after(p_isi, '{{ISI_LAPORAN}}')

new = copy.deepcopy(p_isi._p)
p_lamp._p.addprevious(new)
pl = Paragraph(new, p_lamp._parent)
set_text(pl, 'LAMPIRAN DOKUMENTASI KEGIATAN')
pl.alignment = 1
set_text(p_lamp, '{{LAMPIRAN}}')
p_lamp.runs[0].bold = False
p_lamp.runs[0].underline = False
p_lamp.alignment = 0

d.save(out)
print('Template ditulis:', out)
