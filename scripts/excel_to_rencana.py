#!/usr/bin/env python3
"""
Konversi baseline RBM (Excel) menjadi data rencana untuk form & dashboard.

Pemakaian:
    python scripts/excel_to_rencana.py data/Baseline_RBM_TNTambora_2026.xlsx 2026

Menghasilkan:
    data/rencana_<tahun>.json   -> dibaca form & dashboard
    data/rencana_<tahun>.csv    -> versi tabel untuk dicek manusia
    apps-script/Rencana.gs      -> salinan rencana untuk validasi di server

Kolom Excel yang diharapkan (baris 1 = header, sel bergabung/merge boleh):
    SPTN | Resor | Kode Resor | Indikator Kinerja | Jenis Kegiatan |
    Jumlah Target Kegiatan dalam Satu Tahun | Satuan | Keterangan
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

# Cara menghitung realisasi berdasarkan satuan target.
#   jumlah : jumlahkan kolom "volume" dari laporan yang diterima
#   bulan  : hitung jumlah bulan berbeda yang punya laporan diterima
#   persen : pakai nilai persen dari laporan diterima yang paling akhir
def metode_dari_satuan(satuan: str) -> str:
    s = satuan.strip().lower()
    if s == "bulan":
        return "bulan"
    if s.startswith("persen"):
        return "persen"
    return "jumlah"


def unmerge_value(ws, row, col):
    """Ambil nilai sel; jika bagian dari merge, ambil nilai sel kiri-atas."""
    cell = ws.cell(row=row, column=col)
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            return ws.cell(row=rng.min_row, column=rng.min_col).value
    return cell.value


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/Baseline_RBM_TNTambora_2026.xlsx"
    tahun = int(sys.argv[2]) if len(sys.argv) > 2 else 2026

    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]

    resor_map = {}  # kode pendek -> info resor
    kegiatan = []
    counter = {}

    for r in range(2, ws.max_row + 1):
        nama_keg = clean(ws.cell(row=r, column=5).value)
        target = ws.cell(row=r, column=6).value
        if not nama_keg or target is None:
            continue  # baris kosong / sampah (mis. "\")

        sptn = clean(unmerge_value(ws, r, 1))
        resor = clean(unmerge_value(ws, r, 2))
        kode_resor = clean(unmerge_value(ws, r, 3))
        indikator = clean(unmerge_value(ws, r, 4))
        satuan = clean(ws.cell(row=r, column=7).value)
        ket = clean(ws.cell(row=r, column=8).value)

        rid = kode_resor.split("/")[0]  # "R1/SPTNI/T41" -> "R1"
        if rid not in resor_map:
            resor_map[rid] = {"id": rid, "kode": kode_resor, "nama": resor, "sptn": sptn}

        counter[rid] = counter.get(rid, 0) + 1
        kid = f"{rid}-{counter[rid]:02d}"
        kegiatan.append({
            "id": kid,
            "resor": rid,
            "indikator": indikator,
            "kegiatan": nama_keg,
            "target": float(target) if float(target) != int(target) else int(target),
            "satuan": satuan,
            "metode": metode_dari_satuan(satuan),
            "keterangan": ket,
        })

    indikator_urut = []
    for k in kegiatan:
        if k["indikator"] not in indikator_urut:
            indikator_urut.append(k["indikator"])

    out = {
        "tahun": tahun,
        "sumber": xlsx.name,
        "resor": list(resor_map.values()),
        "indikator": indikator_urut,
        "kegiatan": kegiatan,
    }

    (ROOT / "data").mkdir(exist_ok=True)
    jpath = ROOT / f"data/rencana_{tahun}.json"
    jpath.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")

    cpath = ROOT / f"data/rencana_{tahun}.csv"
    with cpath.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "sptn", "kode_resor", "resor", "indikator", "kegiatan", "target", "satuan", "metode", "keterangan"])
        for k in kegiatan:
            ri = resor_map[k["resor"]]
            w.writerow([k["id"], ri["sptn"], ri["kode"], ri["nama"], k["indikator"], k["kegiatan"],
                        k["target"], k["satuan"], k["metode"], k["keterangan"]])

    # Salinan ringkas untuk Apps Script (validasi kegiatan per resor di server)
    gs_data = {
        "tahun": tahun,
        "resor": {r["id"]: r["nama"] for r in resor_map.values()},
        "kegiatan": {k["id"]: {"resor": k["resor"], "indikator": k["indikator"], "nama": k["kegiatan"], "satuan": k["satuan"], "metode": k["metode"]} for k in kegiatan},
    }
    gpath = ROOT / "apps-script/Rencana.gs"
    gpath.parent.mkdir(exist_ok=True)
    gpath.write_text(
        "// FILE INI DIBUAT OTOMATIS oleh scripts/excel_to_rencana.py — jangan edit manual.\n"
        "// Salin seluruh isinya ke file Rencana.gs di proyek Apps Script.\n"
        f"var RENCANA = {json.dumps(gs_data, ensure_ascii=False, indent=1)};\n",
        encoding="utf-8",
    )

    print(f"{len(resor_map)} resor, {len(kegiatan)} kegiatan -> {jpath.name}, {cpath.name}, Rencana.gs")
    for rid, ri in resor_map.items():
        n = sum(1 for k in kegiatan if k["resor"] == rid)
        print(f"  {rid} {ri['nama']:<20} {ri['sptn']:<22} {n} kegiatan")


if __name__ == "__main__":
    main()
